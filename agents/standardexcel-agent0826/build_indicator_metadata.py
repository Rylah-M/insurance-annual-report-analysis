#!/usr/bin/env python3
"""Build indicator_metadata.xlsx from the plan's initial indicator list."""

from pathlib import Path

import csv
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DICTIONARY_PATH = PROJECT_ROOT / "agents" / "zd-agent0811" / "indicator" / "indicator_dictionary.xlsx"
DATABASE_PATH = PROJECT_ROOT / "database" / "database_result.csv"
OUTPUT_PATH = Path(__file__).resolve().parent / "indicator_metadata.xlsx"

COLUMNS = [
    "indicator_id",
    "indicator_name",
    "indicator_category",
    "standard_unit",
    "analysis_scope",
    "compare_type",
    "direction",
    "horizontal_analysis",
    "vertical_analysis",
    "excel_sheet",
    "description",
]

# indicator_id, indicator_name, indicator_category, analysis_scope, compare_type,
# direction, horizontal_analysis, vertical_analysis, excel_sheet
INITIAL_INDICATORS = [
    ("F001", "净利润", "盈利能力指标", "财险公司", "both", "positive", "是", "是", "产险对标-财务类"),
    ("F002", "承保利润", "盈利能力指标", "财险公司", "both", "positive", "是", "是", "产险对标-财务类"),
    ("F005", "投资收益", "盈利能力指标", "财险公司", "both", "positive", "是", "是", "产险对标-投资类"),
    ("F003", "保险服务收入", "业务规模指标", "财险业务", "both", "positive", "是", "是", "产险对标-业务类"),
    ("B001", "原保险保费收入", "业务规模指标", "财险公司", "both", "positive", "是", "是", "产险对标-业务类"),
    ("B002", "车险保费收入", "业务规模指标", "业务分部", "both", "positive", "是", "是", "产险对标-业务分部"),
    ("B003", "非车险保费收入", "业务规模指标", "业务分部", "both", "positive", "是", "是", "产险对标-业务分部"),
    ("B005", "农业保险保费", "业务规模指标", "业务分部", "both", "positive", "是", "是", "产险对标-业务分部"),
    ("B010", "保费增长率", "业务规模指标", "财险公司", "vertical", "positive", "否", "是", "产险对标-业务类"),
    ("F004", "保险服务费用", "经营效率指标", "财险业务", "both", "negative", "是", "是", "产险对标-经营效率"),
    ("F006", "综合成本率", "经营效率指标", "财险业务", "both", "negative", "是", "是", "产险对标-承保质量"),
    ("F007", "综合赔付率", "经营效率指标", "财险业务", "both", "negative", "是", "是", "产险对标-承保质量"),
    ("F008", "综合费用率", "经营效率指标", "财险业务", "both", "negative", "是", "是", "产险对标-承保质量"),
    ("R001", "核心偿付能力充足率", "风险管理指标", "财险公司", "both", "positive", "是", "是", "产险对标-偿付能力"),
    ("R002", "综合偿付能力充足率", "风险管理指标", "财险公司", "both", "positive", "是", "是", "产险对标-偿付能力"),
]


def load_dictionary():
    wb = openpyxl.load_workbook(DICTIONARY_PATH, data_only=True)
    ws = wb["indicator_dictionary"]
    headers = [str(c.value).strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(headers, row)))
    return {r["indicator_id"]: r for r in rows}


def load_database_rows():
    with DATABASE_PATH.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def main():
    dictionary = load_dictionary()
    db_rows = load_database_rows()

    missing_ids = [ind[0] for ind in INITIAL_INDICATORS if ind[0] not in dictionary]
    if missing_ids:
        raise SystemExit(f"indicator_dictionary.xlsx 中缺少: {missing_ids}")

    eligible = {"财险口径", "特殊财险口径"}
    missing_in_db = []
    for indicator_id, _, _, _, _, _, _, _, _ in INITIAL_INDICATORS:
        has_row = any(
            row["indicator_id"] == indicator_id and row["business_scope_type"] in eligible
            for row in db_rows
        )
        if not has_row:
            missing_in_db.append(indicator_id)
    if missing_in_db:
        raise SystemExit(f"database_result.csv 中缺少财险口径/特殊财险口径数据: {missing_in_db}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "indicator_metadata"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(vertical="top", wrap_text=True)

    ws.append(COLUMNS)
    for col in ws[1]:
        col.fill = header_fill
        col.font = header_font
        col.alignment = center
        col.border = border

    for item in INITIAL_INDICATORS:
        indicator_id = item[0]
        dic = dictionary[indicator_id]
        unit = dic["unit"].strip()
        if "%" in unit:
            standard_unit = "%"
        else:
            standard_unit = unit
        row = list(item)
        row.insert(3, standard_unit)
        row.append(str(dic["definition"]).strip())
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = center if cell.column in {1, 4, 5, 6, 7, 8, 9, 10} else wrap

    widths = [12, 18, 15, 12, 12, 10, 10, 14, 14, 18, 70]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_PATH)
    print(f"已生成: {OUTPUT_PATH}")
    print(f"指标数量: {ws.max_row - 1}")


if __name__ == "__main__":
    main()
