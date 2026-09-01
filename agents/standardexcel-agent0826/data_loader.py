"""Read the three input files used by the standardexcel-agent."""

import csv
from pathlib import Path

import openpyxl


def load_database(path):
    """Load database_result.csv into a list of dictionaries."""
    with Path(path).open(encoding="utf-8-sig", newline="") as fh:
        return [dict(row) for row in csv.DictReader(fh)]


def load_xlsx_sheet(path, sheet_name=None):
    """Load one worksheet as a list of dictionaries keyed by the header row."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    if sheet_name is None:
        sheet_name = workbook.sheetnames[0]
    ws = workbook[sheet_name]
    headers = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue
        row = dict(zip(headers, values))
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def load_indicator_metadata(path):
    """Load indicator_metadata.xlsx (sheet: indicator_metadata)."""
    return load_xlsx_sheet(path, sheet_name="indicator_metadata")


def load_indicator_dictionary(path):
    """Load indicator_dictionary.xlsx (sheet: indicator_dictionary)."""
    return load_xlsx_sheet(path, sheet_name="indicator_dictionary")
