"""Generate standardexcel.xlsx with the all-period peer-comparison layout."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from indicator_filter import parse_indicator_value

TITLE_FONT = Font(name="仿宋", size=16, bold=True)
SUBTITLE_FONT = Font(name="仿宋", size=13, bold=True)
HEADER_FONT = Font(name="仿宋", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="仿宋", size=11)
LABEL_FONT = Font(name="仿宋", size=11, bold=True)
NOTE_FONT = Font(name="仿宋", size=10, color="7F7F7F")
TOTAL_FONT = Font(name="仿宋", size=11, bold=True)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
LABEL_FILL = PatternFill("solid", fgColor="DCE6F1")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
ANGLE_FILL = PatternFill("solid", fgColor="BDD7EE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def _number_format(unit):
    if unit and "%" in str(unit):
        return config.FORMAT_PERCENT
    return config.FORMAT_AMOUNT


def _unit_for_indicator(best_rows, indicator_id, companies, periods, fallback_unit):
    for period in periods:
        for company in companies:
            item = best_rows.get((indicator_id, company, period))
            if item is not None and item[1].get("unit"):
                return item[1]["unit"]
    return fallback_unit


def _period_value(best_rows, indicator_id, company, period):
    item = best_rows.get((indicator_id, company, period))
    if item is None:
        return None
    return parse_indicator_value(item[1].get("indicator_value"))


def _style_header(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER_WRAP
    cell.border = BORDER


def _write_data_row(ws, row, start_col, company, values, unit, total=False):
    font = TOTAL_FONT if total else BODY_FONT
    fill = TOTAL_FILL if total else None

    company_cell = ws.cell(row, start_col, company)
    company_cell.font = font
    company_cell.alignment = CENTER_WRAP
    company_cell.border = BORDER
    if fill:
        company_cell.fill = fill

    for col_idx, value in enumerate(values, start=start_col + 2):
        cell = ws.cell(row, col_idx)
        if value is None:
            cell.value = "-"
        else:
            cell.value = value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = _number_format(unit)
        cell.font = font
        cell.alignment = RIGHT
        cell.border = BORDER
        if fill:
            cell.fill = fill


def _write_table(
    ws, start_row, start_col, metadata_items, best_rows, companies, periods, angle=None
):
    for col_idx, text in enumerate(["公司", "指标"], start=start_col):
        _style_header(ws.cell(start_row, col_idx, text))
    for period_idx, period in enumerate(periods, start=start_col + 2):
        _style_header(ws.cell(start_row, period_idx, period))
    ws.row_dimensions[start_row].height = 24

    current_row = start_row + 1
    for meta in metadata_items:
        indicator_id = meta.get("indicator_id")
        unit = _unit_for_indicator(
            best_rows, indicator_id, companies, periods, meta.get("standard_unit") or ""
        )
        block_start = current_row

        for company in companies:
            values = [_period_value(best_rows, indicator_id, company, period) for period in periods]
            _write_data_row(
                ws,
                current_row,
                start_col,
                config.COMPANY_DISPLAY_NAMES.get(company, company),
                values,
                unit,
            )
            current_row += 1

        totals = []
        is_percent = bool(unit) and "%" in str(unit)
        for period in periods:
            numeric = []
            for company in companies:
                value = _period_value(best_rows, indicator_id, company, period)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric.append(value)
            if not numeric:
                totals.append(None)
            elif is_percent:
                totals.append(sum(numeric) / len(numeric))
            else:
                totals.append(sum(numeric))
        _write_data_row(
            ws, current_row, start_col, config.TOTAL_ROW_LABEL, totals, unit, total=True
        )
        block_end = current_row
        current_row += 1

        for row_num in range(block_start, block_end + 1):
            ws.cell(row_num, start_col + 1).border = BORDER
        indicator_text = f"{meta.get('indicator_name')}（{unit}）"
        indicator_cell = ws.cell(block_start, start_col + 1, indicator_text)
        indicator_cell.font = BODY_FONT
        indicator_cell.alignment = CENTER_WRAP
        indicator_cell.border = BORDER
        ws.merge_cells(
            start_row=block_start,
            start_column=start_col + 1,
            end_row=block_end,
            end_column=start_col + 1,
        )

    end_row = current_row - 1
    if angle is not None:
        ws.merge_cells(
            start_row=start_row,
            start_column=start_col - 1,
            end_row=end_row,
            end_column=start_col - 1,
        )
        angle_cell = ws.cell(start_row, start_col - 1, angle)
        angle_cell.font = LABEL_FONT
        angle_cell.fill = ANGLE_FILL
        angle_cell.alignment = CENTER_WRAP
        angle_cell.border = BORDER
    return start_row, end_row


def _set_table_widths(ws, periods, start_col):
    ws.column_dimensions[get_column_letter(start_col)].width = 18
    ws.column_dimensions[get_column_letter(start_col + 1)].width = 28
    for i in range(len(periods)):
        ws.column_dimensions[get_column_letter(start_col + 2 + i)].width = 18


def _write_comparison_sheet(wb, sheet_name, metadata_items, best_rows, companies, periods):
    ws = wb.create_sheet(sheet_name)
    _write_table(ws, 1, 1, metadata_items, best_rows, companies, periods)
    _set_table_widths(ws, periods, 1)
    ws.freeze_panes = "C2"


def _write_summary_sheet(wb, metadata_items, best_rows, companies, periods):
    ws = wb.create_sheet("汇总")
    row = 1
    for category, sheet_name in config.CATEGORY_SHEETS:
        items = [m for m in metadata_items if (m.get("indicator_category") or "") == category]
        angle = config.CATEGORY_ANGLES[sheet_name]
        _, end = _write_table(ws, row, 2, items, best_rows, companies, periods, angle=angle)
        row = end + 2
    _set_table_widths(ws, periods, 2)
    ws.column_dimensions["A"].width = 16
    ws.freeze_panes = "D2"


def _write_cover_sheet(wb, companies, metadata_count, periods):
    ws = wb.active
    ws.title = "首页说明"

    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = "保险公司年报智能分析Agent"
    title_cell.font = TITLE_FONT
    title_cell.alignment = CENTER_WRAP
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:B2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = "上市财险公司横向对标分析底稿（standardexcel-agent）"
    subtitle_cell.font = SUBTITLE_FONT
    subtitle_cell.alignment = CENTER_WRAP
    ws.row_dimensions[2].height = 26

    company_text = "、".join(config.COMPANY_DISPLAY_NAMES.get(c, c) for c in companies)
    info = [
        ("项目名称", "保险公司年报智能分析Agent"),
        ("数据来源", "上市公司年度报告（database_result.csv，由AI自动提取）"),
        ("数据范围-公司", company_text),
        (
            "数据范围-期间",
            f"{periods[0]} 至 {periods[-1]}（所选年份Q2/Q4）",
        ),
        ("分析指标数", f"{metadata_count}个（规则来自indicator_metadata.xlsx）"),
        ("统计口径", "仅纳入财险口径、特殊财险口径；集团口径指标不参与财险比较"),
    ]
    row = 4
    for label, value in info:
        label_cell = ws.cell(row, 1, label)
        value_cell = ws.cell(row, 2, value)
        label_cell.font = LABEL_FONT
        label_cell.fill = LABEL_FILL
        label_cell.alignment = LEFT_WRAP
        label_cell.border = BORDER
        value_cell.font = BODY_FONT
        value_cell.alignment = LEFT_WRAP
        value_cell.border = BORDER
        ws.row_dimensions[row].height = 24
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    section_cell = ws.cell(row, 1, "说明")
    section_cell.font = SUBTITLE_FONT
    section_cell.alignment = LEFT_WRAP
    row += 1

    notes = [
        "1. 数据由AI自动提取，review_status为“待审核”，使用前建议复核关键指标。",
        "2. business_scope_type用于统一统计口径；集团口径指标不参与财险比较。",
        "3. 指标筛选与分类以indicator_metadata.xlsx为规则中心，指标定义来自indicator_dictionary.xlsx。",
        f"4. 四个对比Sheet按指标分行块：A列为公司，B列为指标名称（含单位），右侧列示{periods[0]}至{periods[-1]}数据。",
        "5. 每个指标块最后一行“所有上市公司合计”：金额类指标取各公司之和，百分比类指标取各公司均值；缺失值以“-”表示。",
        "6. “汇总”Sheet按盈利能力、业务规模、经营效率、偿付能力四个角度纵向合并展示。",
    ]
    for note in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row, 1, note)
        cell.font = BODY_FONT
        cell.alignment = LEFT_WRAP
        ws.row_dimensions[row].height = 22
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    gen_cell = ws.cell(row, 1, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    gen_cell.font = NOTE_FONT
    gen_cell.alignment = LEFT_WRAP

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90


def _build_workbook(metadata_items, best_rows, companies, periods):
    wb = Workbook()
    _write_cover_sheet(wb, companies, len(metadata_items), periods)
    for category, sheet_name in config.CATEGORY_SHEETS:
        items = [m for m in metadata_items if (m.get("indicator_category") or "") == category]
        _write_comparison_sheet(wb, sheet_name, items, best_rows, companies, periods)
    _write_summary_sheet(wb, metadata_items, best_rows, companies, periods)
    return wb


def generate_workbook(output_path, metadata_items, best_rows, companies, periods):
    wb = _build_workbook(metadata_items, best_rows, companies, periods)
    wb.save(output_path)
    return output_path


def generate_workbook_bytes(metadata_items, best_rows, companies, periods):
    wb = _build_workbook(metadata_items, best_rows, companies, periods)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
